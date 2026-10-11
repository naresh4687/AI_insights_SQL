"""Reading uploaded CSV and Excel files (src/ingest/readers.py)."""

import io
import zipfile
from datetime import date, datetime

import openpyxl
import pytest

from src.ingest.readers import ImportFileError, Limits, ReadOptions, read_file

LIMITS = Limits(max_bytes=1_000_000, max_rows=1_000, max_columns=20)


def read_csv(content: bytes | str, **options):
    if isinstance(content, str):
        content = content.encode("utf-8")
    return read_file("data.csv", content, ReadOptions(**options), LIMITS)


def workbook_bytes(sheets: dict[str, list[list]]) -> bytes:
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for title, rows in sheets.items():
        ws = wb.create_sheet(title)
        for row in rows:
            ws.append(row)
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def errors_with(fn, *args, **kwargs) -> str:
    with pytest.raises(ImportFileError) as exc:
        fn(*args, **kwargs)
    return exc.value.message


# --- CSV ------------------------------------------------------------------------

def test_values_are_kept_exactly_as_text():
    raw = read_csv("id,price,code,note\n1, 2.50 ,007,\"a, b\"\n")
    assert raw.headers == ["id", "price", "code", "note"]
    assert raw.frame.values.tolist() == [["1", " 2.50 ", "007", "a, b"]]
    assert (raw.encoding, raw.delimiter, raw.file_type) == ("utf-8", ",", "csv")


@pytest.mark.parametrize("delimiter, char", [("auto", ";"), ("semicolon", ";"), ("tab", "\t"), ("pipe", "|")])
def test_delimiters(delimiter, char):
    raw = read_csv(f"a{char}b\n1{char}2\n3{char}4\n", delimiter=delimiter)
    assert raw.headers == ["a", "b"] and raw.frame.values.tolist() == [["1", "2"], ["3", "4"]]


def test_utf8_bom_is_removed_from_the_first_header():
    raw = read_csv("﻿name,city\nZoë,Chennai\n".encode("utf-8"))
    assert raw.headers == ["name", "city"] and raw.frame.iloc[0, 0] == "Zoë"


def test_utf16_with_bom_is_detected():
    raw = read_csv("name,city\nAnbu,Madurai\n".encode("utf-16"))
    assert raw.encoding == "utf-16" and raw.frame.values.tolist() == [["Anbu", "Madurai"]]


def test_invalid_utf8_is_reported_with_advice_not_guessed():
    content = "name\ncafé\n".encode("cp1252")
    message = errors_with(read_csv, content)
    assert "not UTF-8" in message and "cp1252" in message


def test_explicit_encoding_reads_windows_files():
    raw = read_csv("name\ncafé\n".encode("cp1252"), encoding="cp1252")
    assert raw.frame.iloc[0, 0] == "café"


def test_wrong_explicit_encoding_is_an_error():
    assert "cannot be read as utf-16" in errors_with(read_csv, b"abc,def\n1,2\n3", encoding="utf-16")


@pytest.mark.parametrize("content, expected", [
    (b"", "empty"),
    (b"   \n\n", "empty"),
    (b"a,b\x00c\n1,2\n", "binary"),
    (b'a,b\n1,"unclosed\n', "could not be read"),
])
def test_unreadable_files(content, expected):
    assert expected in errors_with(read_csv, content)


def test_header_only_file_has_zero_rows():
    raw = read_csv("a,b\n")
    assert raw.rows == 0 and raw.headers == ["a", "b"]


def test_ragged_rows_are_recorded_by_file_line():
    raw = read_csv("a,b,c\n1,2,3\n4,5\n\n6,7,8,9\n")
    assert raw.short_rows == [3] and raw.long_rows == [5] and raw.blank_lines == 1
    assert raw.frame.values.tolist() == [["1", "2", "3"], ["4", "5", ""], ["6", "7", "8"]]


def test_duplicate_and_empty_headers_are_kept_by_position():
    raw = read_csv("a,a,,b\n1,2,3,4\n")
    assert raw.headers == ["a", "a", "", "b"] and list(raw.frame.columns) == [0, 1, 2, 3]


def test_no_header_option_names_columns():
    raw = read_csv("1,2\n3,4\n", has_header=False)
    assert raw.headers == ["column_1", "column_2"] and raw.rows == 2


@pytest.mark.parametrize("content, expected", [
    ("a\n" + "1\n" * 1001, "more than 1,000 rows"),
    (",".join(f"c{i}" for i in range(21)) + "\n", "21 columns"),
])
def test_row_and_column_limits(content, expected):
    assert expected in errors_with(read_csv, content)


def test_size_limit():
    assert "larger than the limit" in errors_with(read_file, "big.csv", b"a\n" + b"1\n" * 600_000,
                                                   ReadOptions(), LIMITS)


@pytest.mark.parametrize("name, expected", [("data.xls", ".xls"), ("data.json", "Only CSV"), ("data", "Only CSV"),
                                            ("data.exe", "Only CSV")])
def test_unsupported_file_types(name, expected):
    assert expected in errors_with(read_file, name, b"a\n1\n", ReadOptions(), LIMITS)


# --- Excel ----------------------------------------------------------------------

def test_excel_cells_become_the_text_a_user_sees():
    content = workbook_bytes({"Orders": [
        ["day", "when", "qty", "price", "paid", "note", None],
        [date(2025, 1, 2), datetime(2025, 1, 2, 9, 30), 3, 2.5, True, None, None],
        [datetime(2025, 1, 3), datetime(2025, 1, 3, 0, 0), 4.0, 0.1, False, "x", None],
    ]})
    raw = read_file("orders.xlsx", content, ReadOptions(), LIMITS)
    assert raw.headers == ["day", "when", "qty", "price", "paid", "note"]  # trailing empty column dropped
    assert raw.frame.values.tolist() == [
        ["2025-01-02", "2025-01-02 09:30:00", "3", "2.5", "true", ""],
        ["2025-01-03", "2025-01-03", "4", "0.1", "false", "x"],
    ]
    assert raw.sheets == ["Orders"] and raw.sheet == "Orders"


def test_excel_worksheet_choice():
    content = workbook_bytes({"First": [["a"], [1]], "Second": [["b"], [2], [3]]})
    assert read_file("w.xlsx", content, ReadOptions(), LIMITS).headers == ["a"]
    second = read_file("w.xlsx", content, ReadOptions(sheet="Second"), LIMITS)
    assert second.headers == ["b"] and second.rows == 2 and second.sheets == ["First", "Second"]
    assert "no worksheet named" in errors_with(read_file, "w.xlsx", content, ReadOptions(sheet="Nope"), LIMITS)


def test_excel_row_limit():
    content = workbook_bytes({"S": [["a"]] + [[i] for i in range(1_001)]})
    assert "more than 1,000 rows" in errors_with(read_file, "w.xlsx", content, ReadOptions(), LIMITS)


def test_damaged_and_fake_excel_files():
    assert "not a valid .xlsx" in errors_with(read_file, "w.xlsx", b"not a zip file", ReadOptions(), LIMITS)
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        archive.writestr("hello.txt", "not a workbook")
    assert "could not be read" in errors_with(read_file, "w.xlsx", buffer.getvalue(), ReadOptions(), LIMITS)


def test_excel_zip_bomb_is_refused_before_parsing():
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("xl/worksheets/sheet1.xml", b"0" * (25 * 1_000_000))  # compresses to ~25 KB
    content = buffer.getvalue()
    assert len(content) < LIMITS.max_bytes
    assert "expands to too much data" in errors_with(read_file, "bomb.xlsx", content, ReadOptions(), LIMITS)


def _with_doctype(content: bytes, part: str | None) -> bytes:
    """Rebuild the workbook, putting an entity declaration in front of ``part`` (None = no change)."""
    source = zipfile.ZipFile(io.BytesIO(content))
    doctype = (b'<?xml version="1.0"?><!DOCTYPE x [<!ENTITY a "aaaaaaaaaa">'
               b'<!ENTITY b "&a;&a;&a;&a;&a;&a;&a;&a;&a;&a;">]>')
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as target:
        for name in source.namelist():
            data = source.read(name)
            if name == part:
                assert not data.startswith(b"<?xml")  # the declaration really is planted
                data = doctype + data
            target.writestr(name, data)
    return buffer.getvalue()


def test_rebuilt_workbook_without_attack_still_reads():
    content = _with_doctype(workbook_bytes({"S": [["a"], [1]]}), None)
    assert read_file("ok.xlsx", content, ReadOptions(), LIMITS).frame.values.tolist() == [["1"]]


@pytest.mark.parametrize("part", ["xl/workbook.xml", "xl/worksheets/sheet1.xml"])
def test_excel_xml_entity_attack_is_blocked(part):
    """A 'billion laughs' entity declaration in the workbook XML is refused by defusedxml."""
    content = _with_doctype(workbook_bytes({"S": [["a"], [1]]}), part)
    assert "could not be read" in errors_with(read_file, "evil.xlsx", content, ReadOptions(), LIMITS)
    with pytest.raises(Exception) as exc:  # the refusal comes from defusedxml, not another error
        wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True)
        list(wb.active.iter_rows())
    assert type(exc.value.__cause__ or exc.value.__context__).__name__ == "EntitiesForbidden"
